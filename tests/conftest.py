from pathlib import Path

import pytest
from docx import Document

DATA = Path(__file__).parent / "data" / "policies"


@pytest.fixture
def policies_dir() -> Path:
    return DATA


def _make_pdf(path: Path, pages: list[str]) -> None:
    """Write a tiny valid PDF with one line of text per page (no extra dependencies)."""
    objects: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(n))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + n * 2
    for i, text in enumerate(pages):
        content_id = 4 + i * 2
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(bytes(out))


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    path = tmp_path / "incident-response.pdf"
    _make_pdf(
        path,
        [
            "Security incidents are reported within 72 hours.",
            "A post incident review is held after every major incident.",
        ],
    )
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    path = tmp_path / "vendor-policy.docx"
    doc = Document()
    doc.add_heading("Vendor Management", level=1)
    doc.add_paragraph("All vendors with access to customer data undergo a security review.")
    doc.add_heading("Offboarding", level=1)
    doc.add_paragraph("Vendor access is removed when a contract ends.")
    doc.save(str(path))
    return path
