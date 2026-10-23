"""Find the questions in a customer's questionnaire (XLSX or CSV).

Customer questionnaires have no standard layout, so we detect the header row and the
question/answer columns from common header names, and let the user override the guess
with an explicit column mapping.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from trustreply.core.models import Question, Questionnaire

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".csv"}
HEADER_SCAN_ROWS = 25
MIN_QUESTION_CHARS = 5

_QUESTION_HEADERS = {
    "question",
    "questions",
    "control",
    "controls",
    "control question",
    "requirement",
    "requirements",
    "inquiry",
    "statement",
    "query",
    "control description",
    "question text",
}
_ANSWER_HEADERS = {
    "answer",
    "answers",
    "response",
    "responses",
    "your response",
    "vendor response",
    "supplier response",
    "answer/response",
    "yes/no",
    "compliance",
}
_REFERENCE_HEADERS = {
    "id",
    "ref",
    "reference",
    "no",
    "no.",
    "#",
    "number",
    "question id",
    "control id",
}


class QuestionnaireFormatError(ValueError):
    """Raised when we cannot work out where the questions are."""


_Grid = list[list[str]]


def _norm(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower())


def _load_xlsx(path: Path, sheet: str | None) -> list[tuple[str | None, _Grid]]:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = [sheet] if sheet else workbook.sheetnames
        for name in names:
            if name not in workbook.sheetnames:
                raise QuestionnaireFormatError(
                    f"Sheet '{name}' not found. Available: {', '.join(workbook.sheetnames)}"
                )
        grids: list[tuple[str | None, _Grid]] = []
        for name in names:
            rows = [
                ["" if cell is None else str(cell).strip() for cell in row]
                for row in workbook[name].iter_rows(values_only=True)
            ]
            grids.append((name, rows))
        return grids
    finally:
        workbook.close()


def _load_csv(path: Path) -> list[tuple[str | None, _Grid]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel  # type: ignore[assignment]
    rows = [[cell.strip() for cell in row] for row in csv.reader(text.splitlines(), dialect)]
    return [(None, rows)]


def _find_column(header: list[str], names: set[str]) -> int | None:
    for index, cell in enumerate(header):
        if _norm(cell) in names:
            return index
    return None


def _detect_header(grid: _Grid) -> int | None:
    """Return the 0-based index of the first row that looks like a header row."""
    for index, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        if _find_column(row, _QUESTION_HEADERS) is not None:
            return index
    return None


def _letter_to_index(letter: str) -> int:
    try:
        return column_index_from_string(letter.strip().upper()) - 1
    except ValueError as error:
        raise QuestionnaireFormatError(f"'{letter}' is not a valid column letter.") from error


def _parse_grid(
    grid: _Grid,
    sheet: str | None,
    question_col: str | None,
    answer_col: str | None,
    header_row: int | None,
) -> tuple[list[Question], bool] | None:
    header_index = (header_row - 1) if header_row else _detect_header(grid)

    if question_col:
        q_idx = _letter_to_index(question_col)
        if header_index is None:
            header_index = 0  # explicit mapping without a detectable header: assume row 1
    else:
        if header_index is None:
            return None
        q_idx = _find_column(grid[header_index], _QUESTION_HEADERS)  # type: ignore[assignment]

    header = grid[header_index] if header_index < len(grid) else []
    if answer_col:
        a_idx = _letter_to_index(answer_col)
        answer_exists = True
    else:
        found = _find_column(header, _ANSWER_HEADERS)
        answer_exists = found is not None
        width = max((len(row) for row in grid), default=0)
        a_idx = found if found is not None else max(width, q_idx + 1)
    ref_idx = _find_column(header, _REFERENCE_HEADERS)

    questions: list[Question] = []
    for offset, row in enumerate(grid[header_index + 1 :], start=header_index + 2):
        text = row[q_idx].strip() if q_idx < len(row) else ""
        if len(text) < MIN_QUESTION_CHARS:
            continue
        reference = row[ref_idx].strip() if ref_idx is not None and ref_idx < len(row) else ""
        cell = f"{get_column_letter(q_idx + 1)}{offset}"
        questions.append(
            Question(
                id=f"{sheet}!{cell}" if sheet else cell,
                text=text,
                sheet=sheet,
                row=offset,
                question_col=get_column_letter(q_idx + 1),
                answer_col=get_column_letter(a_idx + 1),
                reference=reference or None,
            )
        )
    return questions, answer_exists


def parse_questionnaire(
    path: str | Path,
    *,
    sheet: str | None = None,
    question_col: str | None = None,
    answer_col: str | None = None,
    header_row: int | None = None,
) -> Questionnaire:
    """Read a questionnaire and list its questions with the cell each answer belongs in.

    Detection uses common header names ("Question", "Control", "Response", ...). If the file
    is unusual, pass ``question_col`` (and optionally ``answer_col``/``header_row``) explicitly.
    Without ``sheet``, every sheet of an XLSX file is scanned and sheets with no questions
    are skipped.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"No such file: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise QuestionnaireFormatError(
            f"Unsupported questionnaire type '{suffix}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    grids = _load_csv(path) if suffix == ".csv" else _load_xlsx(path, sheet)

    questions: list[Question] = []
    answer_exists = True
    for name, grid in grids:
        parsed = _parse_grid(grid, name, question_col, answer_col, header_row)
        if parsed is None:
            continue
        found, exists = parsed
        questions.extend(found)
        answer_exists = answer_exists and exists

    if not questions:
        raise QuestionnaireFormatError(
            "No questions found. Could not detect a header like 'Question' or 'Control'. "
            "Pass the question column explicitly (for example --question-col B)."
        )
    return Questionnaire(source=path.name, questions=questions, answer_column_exists=answer_exists)
